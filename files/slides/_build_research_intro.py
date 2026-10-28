#!/usr/bin/env python3
"""5-slide expert briefing. Palette: premium academic blue."""

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Premium blue
DEEP = RGBColor(0x0A, 0x1F, 0x3D)
PRIMARY = RGBColor(0x12, 0x3A, 0x63)
MID = RGBColor(0x1E, 0x4E, 0x7A)
ACCENT = RGBColor(0x3A, 0x7C, 0xA5)
PALE = RGBColor(0xC5, 0xD6, 0xE6)
WASH = RGBColor(0xE8, 0xF0, 0xF6)
BG = RGBColor(0xF4, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x12, 0x20, 0x33)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
LINE = RGBColor(0xD0, 0xDC, 0xE6)
SOFT = RGBColor(0xB8, 0xCD, 0xDE)

CN = "微软雅黑"
EN = "Calibri"
PT = "/tmp/research-intro-portrait.png"


def rgb_hex(c):
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def set_run_font(run, size, color, bold=False, name=CN):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    latin = rpr.get_or_add_latin()
    latin.set("typeface", name)
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", CN)


def add_box(slide, left, top, width, height, fill=None, line=None, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.shadow.inherit = False
    if radius:
        # 0.05–0.12 keeps a quiet corner; larger values look bubbly.
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Emu(6350)
    return shape


def set_linear_gradient(shape, hex1, hex2, angle=270):
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)
    grad = etree.SubElement(spPr, qn("a:gradFill"))
    gs_lst = etree.SubElement(grad, qn("a:gsLst"))
    for pos, col in ((0, hex1), (100000, hex2)):
        gs = etree.SubElement(gs_lst, qn("a:gs"))
        gs.set("pos", str(pos))
        srgb = etree.SubElement(gs, qn("a:srgbClr"))
        srgb.set("val", col)
    lin = etree.SubElement(grad, qn("a:lin"))
    lin.set("ang", str(int(angle * 60000)))
    lin.set("scaled", "0")


def add_text(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    for i, item in enumerate(lines):
        text, size, color, bold = item[:4]
        space_after = item[4] if len(item) > 4 else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.clear()
        run = p.add_run()
        run.text = text
        set_run_font(run, size, color, bold)
        p.space_after = Pt(space_after)
        p.line_spacing = 1.18
    return box


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def deco_circle(slide, left, top, size, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    # soften by setting transparency on the solid fill
    spPr = shape._element.spPr
    solid = spPr.find(qn("a:solidFill"))
    if solid is not None:
        srgb = solid.find(qn("a:srgbClr"))
        if srgb is not None:
            alpha = etree.SubElement(srgb, qn("a:alpha"))
            alpha.set("val", "18000")  # 18%
    return shape


def footer(slide, page, total=5):
    add_box(slide, Inches(0), Inches(7.26), SLIDE_W, Inches(0.24), fill=DEEP)
    add_text(
        slide,
        Inches(0.5),
        Inches(7.26),
        Inches(10.2),
        Inches(0.24),
        [("徐安察    浙江工商大学统计与数据科学学院", 10, SOFT, False)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        Inches(11.5),
        Inches(7.26),
        Inches(1.4),
        Inches(0.24),
        [(f"{page:02d}  /  {total:02d}", 10, PALE, False)],
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def header(slide, kicker, title):
    bar = add_box(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), fill=DEEP)
    set_linear_gradient(bar, "0A1F3D", "1E4E7A", angle=0)
    add_text(slide, Inches(0.55), Inches(0.26), Inches(12.2), Inches(0.26), [(kicker, 12, ACCENT, False)])
    add_text(slide, Inches(0.55), Inches(0.52), Inches(12.2), Inches(0.44), [(title, 24, DEEP, True)])
    add_box(slide, Inches(0.55), Inches(1.02), Inches(0.42), Inches(0.04), fill=ACCENT)


def make_portrait():
    im = Image.open("/workspace/images/ancha.png").convert("RGB")
    w, h = im.size
    side = w
    top = max(int(h * 0.02), 0)
    if top + side > h:
        top = max(h - side, 0)
    face = im.crop((0, top, w, top + side)).resize((900, 900), Image.Resampling.LANCZOS)
    bg = Image.new("RGB", (900, 900), (10, 31, 61))
    mask = Image.new("L", (900, 900), 0)
    ImageDraw.Draw(mask).ellipse((10, 10, 889, 889), fill=255)
    bg.paste(face, (0, 0), mask)
    ImageDraw.Draw(bg).ellipse((8, 8, 891, 891), outline=(58, 124, 165), width=12)
    bg.save(PT)


def build():
    make_portrait()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ------------------------------------------------------------------
    # 1  个人研究与业绩
    # ------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    set_bg(s, BG)
    rail = add_box(s, Inches(0), Inches(0), Inches(4.05), SLIDE_H, fill=DEEP)
    set_linear_gradient(rail, "0A1F3D", "123A63", angle=200)
    deco_circle(s, Inches(-1.1), Inches(-1.0), Inches(3.2), RGBColor(0x3A, 0x7C, 0xA5))
    deco_circle(s, Inches(2.3), Inches(5.6), Inches(2.6), RGBColor(0x1E, 0x4E, 0x7A))

    s.shapes.add_picture(PT, Inches(0.78), Inches(0.55), Inches(2.48), Inches(2.48))
    add_text(s, Inches(0.25), Inches(3.22), Inches(3.55), Inches(0.48), [("徐安察", 28, WHITE, True)], align=PP_ALIGN.CENTER)
    add_text(
        s,
        Inches(0.25),
        Inches(3.72),
        Inches(3.55),
        Inches(0.32),
        [("教授  ·  博士生导师", 13, ACCENT, True)],
        align=PP_ALIGN.CENTER,
    )
    add_box(s, Inches(1.70), Inches(4.14), Inches(0.64), Inches(0.03), fill=ACCENT)
    add_text(
        s,
        Inches(0.25),
        Inches(4.32),
        Inches(3.55),
        Inches(1.15),
        [
            ("浙江工商大学", 13, PALE, False, 3),
            ("统计与数据科学学院", 13, PALE, False, 10),
            ("xuancha@mail.zjgsu.edu.cn", 11, SOFT, False, 3),
            ("www.xu-ancha.com", 11, SOFT, False),
        ],
        align=PP_ALIGN.CENTER,
    )

    add_text(s, Inches(4.45), Inches(0.36), Inches(8.4), Inches(0.24), [("浙江工商大学统计与数据科学学院", 12, ACCENT, False)])
    add_text(s, Inches(4.45), Inches(0.62), Inches(8.4), Inches(0.44), [("个人简介", 28, DEEP, True)])
    add_box(s, Inches(4.45), Inches(1.12), Inches(0.42), Inches(0.04), fill=ACCENT)
    add_text(
        s,
        Inches(4.45),
        Inches(1.28),
        Inches(8.4),
        Inches(0.55),
        [
            ("主要做退化数据、寿命数据和贝叶斯方法。", 15, INK, False, 2),
            ("下面分教学、人才培养和科研说一下。", 15, INK, False),
        ],
    )

    add_box(s, Inches(4.45), Inches(1.98), Inches(4.20), Inches(4.92), fill=WHITE, line=LINE, radius=0.05)
    add_box(s, Inches(4.45), Inches(1.98), Inches(0.07), Inches(4.92), fill=ACCENT)
    add_text(s, Inches(4.72), Inches(2.12), Inches(3.75), Inches(0.32), [("教学与人才培养", 15, DEEP, True)])
    add_text(
        s,
        Inches(4.72),
        Inches(2.52),
        Inches(3.75),
        Inches(4.20),
        [
            ("本科开过统计质量管理、贝叶斯数据", 13, INK, False, 2),
            ("分析、概率论；研究生开过可靠性、", 13, INK, False, 2),
            ("贝叶斯计算。", 13, INK, False, 10),
            ("带博士和硕士。学生里有人拿到国家", 13, INK, False, 2),
            ("奖学金，也有人去了南京理工大学、", 13, INK, False, 2),
            ("温州大学等高校。", 13, INK, False, 10),
            ("两次评上「我心目中的好导师」。", 13, INK, False, 10),
            ("带过挑战杯，作品拿到过全国二等奖。", 13, INK, False),
        ],
    )

    add_box(s, Inches(8.85), Inches(1.98), Inches(4.00), Inches(4.92), fill=WHITE, line=LINE, radius=0.05)
    add_box(s, Inches(8.85), Inches(1.98), Inches(0.07), Inches(4.92), fill=PRIMARY)
    add_text(s, Inches(9.12), Inches(2.12), Inches(3.55), Inches(0.32), [("科学研究", 15, DEEP, True)])
    add_text(
        s,
        Inches(9.12),
        Inches(2.52),
        Inches(3.55),
        Inches(4.20),
        [
            ("一作或通讯六十多篇，其中 6 篇是", 13, INK, False, 2),
            ("ESI 高被引。刊物主要是 NRL、SJS、", 13, INK, False, 2),
            ("JQT、EJOR、IEEE TR、IISE。", 13, INK, False, 10),
            ("2025 年出了一本《随机退化过程与", 13, INK, False, 2),
            ("统计建模》。", 13, INK, False, 10),
            ("主持国自然青年基金和三项面上，", 13, INK, False, 2),
            ("还有省基金重点。", 13, INK, False, 10),
            ("省自然科学奖。现在是可靠性分会", 13, INK, False, 2),
            ("副理事长，STARF 副主编。", 13, INK, False),
        ],
    )
    add_notes(
        s,
        "先报姓名和单位。教学点课程和学生去向，提一句好导师和挑战杯即可。"
        "科研报论文、书、基金和分会兼职。不要念完每一行。",
    )

    # ------------------------------------------------------------------
    # 2  研究背景与定位
    # ------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    set_bg(s)
    header(s, "背景", "研究背景")
    footer(s, 2)

    add_box(s, Inches(0.50), Inches(1.24), Inches(12.30), Inches(2.35), fill=WHITE, line=LINE, radius=0.04)
    add_box(s, Inches(0.50), Inches(1.24), Inches(0.07), Inches(2.35), fill=ACCENT)
    add_text(
        s,
        Inches(0.78),
        Inches(1.40),
        Inches(11.80),
        Inches(2.05),
        [
            ("很多高可靠产品，试验期内几乎看不到失效。寿命数据大多是删失，能用的信息很少。", 15, INK, False, 8),
            ("现在更常见的是把电流、裂纹、光强、转矩这些量跟着测下来。轨迹本身比失效时间细。", 15, INK, False, 8),
            ("问题也跟着来：样本往往不大，指标常常不止一个，个体差异明显，数据还是一批批到的。", 15, INK, False, 8),
            ("先验稍微定偏，后验就会跟着走；每来一批新数据就重估一遍，也不方便。", 15, INK, False),
        ],
    )

    add_text(s, Inches(0.50), Inches(3.76), Inches(12.3), Inches(0.30), [("做的时候会碰到的几个具体问题", 14, DEEP, True)])
    constraints = [
        ("样本小", "失效少，先验不好定"),
        ("指标多", "几个量一起退化，还可能相关"),
        ("个体差得大", "有随机效应，也有两阶段、两种尺度"),
        ("数据陆续来", "不太适合每次重跑全样本"),
    ]
    for i, (t, d) in enumerate(constraints):
        x = Inches(0.50) + i * Inches(3.18)
        add_box(s, x, Inches(4.14), Inches(3.05), Inches(1.18), fill=WHITE, line=LINE, radius=0.05)
        add_text(s, x + Inches(0.16), Inches(4.24), Inches(2.75), Inches(0.32), [(t, 15, DEEP, True)])
        add_text(s, x + Inches(0.16), Inches(4.58), Inches(2.75), Inches(0.60), [(d, 13, MUTED, False)])

    band = add_box(s, Inches(0.50), Inches(5.50), Inches(12.30), Inches(1.54), fill=DEEP, radius=0.04)
    set_linear_gradient(band, "0A1F3D", "1E4E7A", angle=0)
    add_text(s, Inches(0.72), Inches(5.62), Inches(11.90), Inches(0.26), [("我们现在主要做的", 12, ACCENT, False)])
    add_text(
        s,
        Inches(0.72),
        Inches(5.92),
        Inches(11.90),
        Inches(0.95),
        [
            ("把退化过程写清楚，再用贝叶斯方法做可靠性评估和剩余寿命预测。", 15, WHITE, False, 6),
            ("具体包括过程建模、在线更新、小样本下的客观先验，以及加速试验和现场数据怎么放在一起。", 14, PALE, False),
        ],
    )
    add_notes(
        s,
        "先把背景讲完：失效少，所以看退化轨迹。再提四个具体麻烦。最后两句说我们做什么，不要念成口号。",
    )

    # ------------------------------------------------------------------
    # 3  过程建模
    # ------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    set_bg(s)
    header(s, "过程", "退化过程上的几项工作")
    footer(s, 3)

    add_text(
        s,
        Inches(0.55),
        Inches(1.20),
        Inches(12.20),
        Inches(0.50),
        [
            ("还是 Wiener、逆高斯、Gamma 这几类。后来多半是在这上面处理指标相关、两阶段、两种时间尺度，以及尾部偏厚的情形。", 14, MUTED, False),
        ],
    )

    models = [
        ("多元 Wiener", "二元和层次模型，退化速率和波动可以相关", "IEEE TR，2018、2026"),
        ("多元 Student-t", "用来拟合尾部比较重的多指标退化", "IISE Transactions，2025"),
        ("两阶段逆高斯", "重参数化以后做变点和剩余寿命", "EJOR，2024"),
        ("双尺度逆高斯", "带随机效应，两种时间尺度可以放在一起", "IISE Transactions，Featured Article"),
    ]
    for i, (t, d, cite) in enumerate(models):
        col, row = i % 2, i // 2
        x = Inches(0.50) + col * Inches(6.40)
        y = Inches(1.80) + row * Inches(1.72)
        add_box(s, x, y, Inches(6.20), Inches(1.58), fill=WHITE, line=LINE, radius=0.05)
        add_box(s, x, y, Inches(0.07), Inches(1.58), fill=ACCENT if row == 0 else PRIMARY)
        add_text(s, x + Inches(0.28), y + Inches(0.16), Inches(5.70), Inches(0.32), [(t, 16, DEEP, True)])
        add_text(s, x + Inches(0.28), y + Inches(0.52), Inches(5.70), Inches(0.50), [(d, 13, INK, False)])
        add_text(s, x + Inches(0.28), y + Inches(1.10), Inches(5.70), Inches(0.28), [(cite, 11, ACCENT, False)])

    band = add_box(s, Inches(0.50), Inches(5.36), Inches(12.30), Inches(1.66), fill=DEEP, radius=0.04)
    set_linear_gradient(band, "0A1F3D", "1E4E7A", angle=0)
    add_text(s, Inches(0.74), Inches(5.50), Inches(11.85), Inches(0.24), [("书和在做的项目", 12, ACCENT, False)])
    add_text(
        s,
        Inches(0.74),
        Inches(5.78),
        Inches(11.85),
        Inches(1.00),
        [
            ("去年出了一本《随机退化过程与统计建模》，和汤银才、庄亮亮一起写的。", 14, WHITE, False, 4),
            ("手上的面上项目是多维偏态厚尾退化，以及相应的贝叶斯试验设计。", 14, PALE, False),
        ],
    )
    add_notes(
        s,
        "点两篇就够：Student-t 和两阶段逆高斯。书一句话带过。",
    )

    # ------------------------------------------------------------------
    # 4  在线推断
    # ------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    set_bg(s)
    header(s, "推断", "在线推断和剩余寿命")
    footer(s, 4)

    add_text(
        s,
        Inches(0.55),
        Inches(1.20),
        Inches(12.20),
        Inches(0.48),
        [("退化监测一般是陆续来的。后验如果能递推，就不必每来一批数据重新估一遍。Gamma 和逆高斯在合适的参数化下可以走共轭。", 14, MUTED, False)],
    )

    steps = [
        ("共轭", "先验选得好，后验还能写成同类"),
        ("递推", "新数据进来，更新几个超参数"),
        ("剩余寿命", "用当前后验直接算，不必重估"),
        ("异常点", "个别离谱的观测，不要把整条轨迹带偏"),
    ]
    for i, (t, d) in enumerate(steps):
        x = Inches(0.50) + i * Inches(3.18)
        dark = i in (1, 2)
        add_box(s, x, Inches(1.80), Inches(3.05), Inches(1.48), fill=DEEP if dark else WHITE, line=None if dark else LINE, radius=0.05)
        add_text(s, x + Inches(0.16), Inches(1.94), Inches(2.75), Inches(0.34), [(t, 16, WHITE if dark else DEEP, True)])
        add_text(s, x + Inches(0.16), Inches(2.34), Inches(2.75), Inches(0.76), [(d, 13, PALE if dark else MUTED, False)])

    papers = [
        ("Gamma 过程", "共轭先验下的递归预测", "Scandinavian Journal of Statistics，2026"),
        ("逆高斯过程", "在线估计和剩余寿命", "Naval Research Logistics，2025"),
        ("维修决策", "把剩余寿命用到维修安排上", "RESS，2023"),
        ("航空发动机", "数据不出域的联邦学习", "Applied Soft Computing，2025"),
    ]
    for i, (t, d, cite) in enumerate(papers):
        col, row = i % 2, i // 2
        x = Inches(0.50) + col * Inches(6.40)
        y = Inches(3.70) + row * Inches(1.58)
        add_box(s, x, y, Inches(6.20), Inches(1.44), fill=WHITE, line=LINE, radius=0.05)
        add_text(s, x + Inches(0.22), y + Inches(0.14), Inches(5.80), Inches(0.30), [(t, 15, DEEP, True)])
        add_text(s, x + Inches(0.22), y + Inches(0.48), Inches(5.80), Inches(0.40), [(d, 13, INK, False)])
        add_text(s, x + Inches(0.22), y + Inches(0.96), Inches(5.80), Inches(0.28), [(cite, 11, ACCENT, False)])

    add_notes(
        s,
        "统计同行多讲 SJS 和 NRL。工业上可以提航空发动机。不要说“不是更快的 MCMC”这种对比句。",
    )

    # ------------------------------------------------------------------
    # 5  评估与讨论
    # ------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    set_bg(s)
    header(s, "其他", "小样本、应用，还有没做完的")
    footer(s, 5)

    cols = [
        (
            ACCENT,
            "小样本",
            [
                "博士阶段开始做参考先验和后验正常性",
                "后来做到屏蔽数据、竞争失效",
                "也做加速寿命和加速退化",
                "主要写在 IEEE TR、CSDA、JSPI",
            ],
        ),
        (
            PRIMARY,
            "应用",
            [
                "永磁制动器的小样本评估",
                "和华为做过 Weibull 算法、隐患预警",
                "动态环境下的系统可靠性",
                "加速试验和现场数据怎么一起用",
            ],
        ),
        (
            DEEP,
            "后面还想搞清楚的",
            [
                "多维偏态厚尾，参数怎么定才稳",
                "在线更新之后，区间还准不准",
                "试验往现场外推，偏差从哪来",
                "试验设计怎么同时照顾评估和预测",
            ],
        ),
    ]
    for i, (bar, title, bullets) in enumerate(cols):
        x = Inches(0.50) + i * Inches(4.20)
        add_box(s, x, Inches(1.24), Inches(4.02), Inches(4.00), fill=WHITE, line=LINE, radius=0.05)
        add_box(s, x, Inches(1.24), Inches(4.02), Inches(0.06), fill=bar)
        add_text(s, x + Inches(0.20), Inches(1.46), Inches(3.62), Inches(0.50), [(title, 15, DEEP, True)])
        lines = [(f"·  {b}", 13, INK, False, 8) for b in bullets]
        add_text(s, x + Inches(0.20), Inches(2.06), Inches(3.62), Inches(2.95), lines)

    band = add_box(s, Inches(0.50), Inches(5.42), Inches(12.30), Inches(1.62), fill=DEEP, radius=0.04)
    set_linear_gradient(band, "0A1F3D", "1E4E7A", angle=0)
    add_text(
        s,
        Inches(0.74),
        Inches(5.60),
        Inches(8.4),
        Inches(1.20),
        [
            ("后面这几件事还没有想清楚，希望听听各位的意见。", 15, WHITE, False, 6),
            ("谢谢。", 14, PALE, False),
        ],
    )
    add_text(
        s,
        Inches(9.30),
        Inches(5.66),
        Inches(3.25),
        Inches(1.15),
        [
            ("xuancha@mail.zjgsu.edu.cn", 13, WHITE, False, 4),
            ("www.xu-ancha.com", 13, PALE, False),
        ],
    )
    add_notes(
        s,
        "小样本举屏蔽数据或加速试验，应用举永磁制动器或华为。后面的问题挑一两个问在座的人。不要再报论文篇数。",
    )

    out = "/workspace/files/slides/research-intro-experts.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
